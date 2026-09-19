# backend/tests/test_document_docx_inplace.py
"""docx 整篇原位编辑:未变段落零损伤、变更/删除/插入、结构漂移降级。"""

import os
import sys

BACKEND_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if BACKEND_ROOT not in sys.path:
    sys.path.insert(0, BACKEND_ROOT)
TEST_ROOT = os.path.abspath(os.path.dirname(__file__))
if TEST_ROOT not in sys.path:
    sys.path.insert(0, TEST_ROOT)

import copy  # noqa: E402

from docx import Document as DocxDocument  # noqa: E402

from app.domains.asset.services.document.docx import (  # noqa: E402
    apply_blocks_to_docx_inplace,
)
from app.domains.asset.services.document.payload import parse_document_payload  # noqa: E402


def _build_base_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Requirement Overview", level=1)
    p = doc.add_paragraph()
    bold = p.add_run("Login must work.")
    bold.bold = True
    doc.add_paragraph("Session expires in 30 minutes.")
    buf = __import__("io").BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _write_tmp(data: bytes) -> str:
    import tempfile

    fd, path = tempfile.mkstemp(suffix=".docx")
    with os.fdopen(fd, "wb") as f:
        f.write(data)
    return path


def _base_blocks(docx_bytes: bytes):
    payload = parse_document_payload("spec.docx", docx_bytes)
    return list(payload.get("blocks_json") or [])


def _find_block(blocks, needle: str):
    return next(b for b in blocks if needle in str(b.get("text") or ""))


def test_unchanged_blocks_keep_bold_run_and_changed_text_applies():
    original = _build_base_docx_bytes()
    blocks = _base_blocks(original)
    final = copy.deepcopy(blocks)
    target = _find_block(final, "Session expires")
    target["text"] = "Session expires in 60 minutes."
    target["runs"] = [{"text": "Session expires in 60 minutes."}]

    out = apply_blocks_to_docx_inplace(_write_tmp(original), blocks, final)
    assert out is not None
    reparsed = _base_blocks(out)
    assert any("Session expires in 60 minutes." == str(b.get("text") or "") for b in reparsed)
    # 未变段落:粗体 run 原样保留
    login = next(b for b in reparsed if "Login must work." in str(b.get("text") or ""))
    assert any(run.get("bold") for run in (login.get("runs") or []))


def test_deleted_base_block_removes_paragraph():
    original = _build_base_docx_bytes()
    blocks = _base_blocks(original)
    final = [b for b in copy.deepcopy(blocks) if "Session expires" not in str(b.get("text") or "")]

    out = apply_blocks_to_docx_inplace(_write_tmp(original), blocks, final)
    assert out is not None
    reparsed = _base_blocks(out)
    assert all("Session expires" not in str(b.get("text") or "") for b in reparsed)


def test_inserted_new_block_lands_in_order():
    original = _build_base_docx_bytes()
    blocks = _base_blocks(original)
    final = copy.deepcopy(blocks)
    insert_at = next(i for i, b in enumerate(final) if "Login must work." in str(b.get("text") or "")) + 1
    final.insert(
        insert_at,
        {
            "id": "blk-new-1",
            "type": "paragraph",
            "text": "Inserted requirement.",
            "runs": [{"text": "Inserted requirement."}],
        },
    )

    out = apply_blocks_to_docx_inplace(_write_tmp(original), blocks, final)
    assert out is not None
    reparsed = _base_blocks(out)
    texts = [str(b.get("text") or "") for b in reparsed]
    assert "Inserted requirement." in texts
    assert texts.index("Inserted requirement.") == texts.index(
        next(t for t in texts if "Login must work." in t)
    ) + 1


def test_structure_drift_returns_none():
    original = _build_base_docx_bytes()
    blocks = _base_blocks(original)
    # 模拟元素数与 base blocks 对不上:传入空 final 触发元素数量不匹配防御
    out = apply_blocks_to_docx_inplace(_write_tmp(original), blocks, [])
    assert out is None


def test_apply_resolution_document_scope_uses_inplace_and_preserves_bold(tmp_path):
    from app.domains.asset.models.asset import AssetType, SddAsset, SddAssetVersion
    from app.domains.asset.services import asset_resolution_service
    from app.domains.asset.services.asset_discussion_service import create_thread
    from test_workspace_asset_boundary import _build_db, _seed_workspace, _session

    original = _build_base_docx_bytes()
    engine, SessionLocal = _build_db()
    try:
        with _session(SessionLocal) as db:
            _user, _ws, task = _seed_workspace(db)
            task.project_path = str(tmp_path)
            db.commit()

            asset = SddAsset(
                task_id=task.id, workspace_id=task.workspace_id, creator_id=task.creator_id,
                asset_type=AssetType.SPEC, name="spec.docx",
                source_file_name="spec.docx", source_ext=".docx",
                source_mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            db.add(asset)
            db.flush()
            path = os.path.join(str(tmp_path), "v1.docx")
            with open(path, "wb") as f:
                f.write(original)
            version = SddAssetVersion(
                asset_id=asset.id, version_no=1, original_path=path,
                original_ext=".docx", original_mime=asset.source_mime,
                normalized_markdown="", blocks_json=_base_blocks(original),
                render_json={"format": "rich_doc", "block_count": 3},
                created_by=task.creator_id,
            )
            db.add(version)
            db.commit()
            asset.active_version_id = version.id
            db.commit()

            blocks = _base_blocks(original)
            target = _find_block(blocks, "Session expires")
            thread = create_thread(
                db, asset=asset, version=version, creator_id=task.creator_id,
                block_id=str(target["id"]), body="请更新会话时长",
            )
            proposal = asset_resolution_service.create_resolution_proposal(
                db, thread=thread, creator_id=task.creator_id,
                proposed_text="改为 60 分钟", version=version,
            )
            final_blocks = copy.deepcopy(blocks)
            tgt = _find_block(final_blocks, "Session expires")
            tgt["text"] = "Session expires in 60 minutes."
            tgt["runs"] = [{"text": "Session expires in 60 minutes."}]

            new_version = asset_resolution_service.apply_resolution_proposal(
                db, asset=asset, thread=thread, proposal=proposal,
                actor_user_id=task.creator_id,
                final_block_ast=None, final_blocks_ast=final_blocks,
                change_note="test",
            )
            with open(new_version.original_path, "rb") as f:
                out_bytes = f.read()
            reparsed = _base_blocks(out_bytes)
            assert any("Session expires in 60 minutes." == str(b.get("text") or "") for b in reparsed)
            login = next(b for b in reparsed if "Login must work." in str(b.get("text") or ""))
            assert any(run.get("bold") for run in (login.get("runs") or []))
    finally:
        engine.dispose()
